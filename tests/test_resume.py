import sys
from pathlib import Path

import pytest
from weasyprint import HTML

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from werkzeug.datastructures import MultiDict  # noqa: E402

from app import app, download_name  # noqa: E402
from render import LocalStaticFetcher, Style, available_templates, render_html, render_page_images  # noqa: E402
from render_docx import DOCX_MIME, render_docx  # noqa: E402
from resume import Resume, Section, from_form, parse_markdown, validate  # noqa: E402

SAMPLE = (ROOT / "samples" / "sample.md").read_text(encoding="utf-8")


@pytest.fixture
def client():
    app.config["TESTING"] = True
    return app.test_client()


# Markdown parsing

def test_sample_name_headline_contacts():
    r = parse_markdown(SAMPLE)
    assert r.name == "Nurul Aisyah Kamarudin"
    assert r.headline == "Software Engineer"
    assert r.contacts[0] == "aisyah.kamarudin@gmail.com"
    assert "Cyberjaya, Selangor" in r.contacts
    assert len(r.contacts) == 5


def test_sample_sections_in_order():
    r = parse_markdown(SAMPLE)
    assert [s.title for s in r.sections] == [
        "Summary", "Experience", "Education", "Projects", "Skills", "Certifications",
    ]
    experience = r.sections[1].html
    assert experience.count('<div class="entry">') == 2
    assert "<h3>Software Engineering Intern</h3>" in experience


def test_markdown_italic_line_becomes_right_date():
    html = parse_markdown(SAMPLE).sections[1].html
    assert '<h3>Software Engineering Intern</h3><span class="entry-date">Jun 2025 to Sep 2025</span></div>' in html
    assert "<p>Tanjung Data Sdn Bhd, Kuala Lumpur</p>" in html
    # A bold line is not mistaken for a date.
    html = parse_markdown("# A\n\n## X\n### Role\n**Big win**").sections[0].html
    assert "entry-date" not in html and "<strong>Big win</strong>" in html


def test_contacts_without_headline():
    r = parse_markdown("# Ali\nali@mail.com | 012-345 6789\n\n## Skills\n- Python")
    assert r.headline == ""
    assert r.contacts == ["ali@mail.com", "012-345 6789"]


def test_empty_markdown_sections_dropped():
    r = parse_markdown("# Ali\n\n## Summary\n\n## Skills\n- Python\n## Projects\n   \n")
    assert [s.title for s in r.sections] == ["Skills"]


def test_empty_form_sections_dropped():
    form = MultiDict([
        ("name", "Ali"), ("summary", "  "), ("email", "ali@mail.com"), ("links", "a.com, b.com"),
        ("experience_title", ""), ("experience_org", ""), ("experience_details", ""),
        ("skills_label", ""), ("skills_items", "Python, SQL"),
    ])
    r = from_form(form)
    assert [s.title for s in r.sections] == ["Skills"]
    assert r.sections[0].html == "<ul><li>Python, SQL</li></ul>"
    assert r.contacts == ["ali@mail.com", "a.com", "b.com"]


def test_form_entries_build_structure_for_the_user():
    form = MultiDict([
        ("name", "Ali"),
        ("experience_title", "Intern"), ("experience_org", "Tanjung Data"), ("experience_location", "KL"),
        ("experience_dates", "2025"), ("experience_details", "- Built things\n\n• Fixed **bugs**\nShipped"),
        ("experience_title", "Tutor"), ("experience_org", ""), ("experience_location", ""),
        ("experience_dates", ""), ("experience_details", ""),
        ("skills_label", "Languages"), ("skills_items", "Python"),
    ])
    r = from_form(form)
    exp = r.sections[0].html
    assert exp.count('<div class="entry">') == 2
    assert '<h3>Intern</h3><span class="entry-date">2025</span>' in exp
    assert "<p>Tanjung Data, KL</p>" in exp
    assert "<li>Built things</li><li>Fixed <strong>bugs</strong></li><li>Shipped</li>" in exp
    assert '<div class="entry-head"><h3>Tutor</h3></div>' in exp
    assert r.sections[1].html == "<ul><li><strong>Languages:</strong> Python</li></ul>"


def test_details_keep_years_and_strip_real_list_markers():
    r = from_form(MultiDict([("name", "A"), ("experience_title", "T"),
                             ("experience_details", "2024. Won the hackathon\n12.5% growth\n1) First\n- Second")]))
    assert "<li>2024. Won the hackathon</li><li>12.5% growth</li><li>First</li><li>Second</li>" in r.sections[0].html


def test_stray_hash_lines_do_not_become_headings():
    r = from_form(MultiDict([("name", "A"), ("summary", "# Big\nHello")]))
    assert "<h1>" not in r.sections[0].html
    r = parse_markdown("# Ali\n\n## Skills\n# Oops\nSetext\n---\n- Python")
    assert "<h1>" not in r.sections[0].html and "<h2>" not in r.sections[0].html


def test_entry_with_dates_but_no_title_has_no_empty_heading():
    r = from_form(MultiDict([("name", "A"), ("experience_dates", "2024"), ("experience_org", "Co")]))
    assert r.sections[0].html == '<div class="entry"><div class="entry-head"><span class="entry-date">2024</span></div><p>Co</p></div>'


def test_validation_messages():
    assert "markdown" in validate(parse_markdown("no heading here"), "markdown")
    assert set(validate(from_form({}), "form")) == {"name", "sections"}


# Safety

def test_raw_html_is_escaped():
    r = parse_markdown('# <b>Ali</b>\n\n## Skills\n- <script>alert(1)</script>\n- <img src=x onerror=alert(1)>')
    html = render_html(r, "classic")
    assert "<script>" not in html
    assert "&lt;script&gt;" in html
    assert "<img" not in html
    assert "&lt;b&gt;Ali&lt;/b&gt;" in html


def test_form_fields_escaped():
    r = from_form(MultiDict([("name", "A"), ("experience_title", "<script>x</script>"), ("experience_details", "<img src=x>")]))
    html = render_html(r, "modern")
    assert "<script>x" not in html and "<img" not in html
    assert "&lt;script&gt;x&lt;/script&gt;" in html


def test_style_whitelist():
    style = Style.from_form({"accent": "red; } body { display:none", "header_align": "center", "dates": "sideways"})
    assert style == Style(header_align="center")
    assert "display:none" not in style.css()
    css = Style(accent="navy", title_align="center", body_align="justify", dates="below").css()
    assert "--accent: #1f3a5f" in css
    assert ".section-title { text-align: center; }" in css
    assert "text-align: justify" in css
    assert ".entry-head { display: block; }" in css
    assert Style().css() == ""


def test_markdown_images_not_rendered():
    r = parse_markdown("# Ali\n\n## Skills\n![x](http://169.254.169.254/latest/meta-data)")
    assert "<img" not in render_html(r, "classic")


def test_javascript_links_dropped():
    r = parse_markdown("# Ali\n\n## Links\n[site](javascript:alert(1))")
    assert "href" not in r.sections[0].html


def test_unknown_template_rejected():
    r = Resume(name="Ali", sections=[Section("Skills", "<p>x</p>")])
    for bad in ["nope", "../templates/index", "classic/../../app.py", ""]:
        with pytest.raises(ValueError):
            render_html(r, bad)
    assert available_templates() == ["classic", "minimal", "modern"]


def test_fetcher_blocks_remote_and_outside_files():
    fetcher = LocalStaticFetcher()
    for url in ["http://example.com/x.png", "https://169.254.169.254/", (ROOT / "app.py").as_uri(), "file:///etc/passwd"]:
        with pytest.raises(ValueError):
            fetcher.fetch(url)
    font = ROOT / "static" / "fonts" / "geist-400-normal.woff2"
    assert fetcher.fetch(font.as_uri()) is not None


# Routes

def test_index_shows_form(client):
    res = client.get("/")
    assert res.status_code == 200
    assert b'name="markdown"' in res.data and b'name="template"' in res.data


def test_preview_renders_resume(client):
    res = client.post("/preview", data={"mode": "markdown", "markdown": SAMPLE, "template": "modern", "accent": "burgundy"})
    assert res.status_code == 200
    assert b"Nurul Aisyah Kamarudin" in res.data
    assert b'src="data:image/webp;base64,' in res.data
    assert b"Page 1 of 1" in res.data
    assert b"page-warning" not in res.data
    # Settings are controls on the page, not duplicated as hidden fields.
    assert b'type="hidden" name="accent"' not in res.data
    assert b'type="hidden" name="markdown"' in res.data


def test_preview_keeps_repeated_entries(client):
    data = MultiDict([("mode", "form"), ("name", "Ali"), ("experience_title", "One"), ("experience_title", "Two")])
    res = client.post("/preview", data=data)
    assert res.data.count(b'type="hidden" name="experience_title"') == 2


def test_pdf_download(client):
    res = client.post("/pdf", data={"mode": "markdown", "markdown": SAMPLE, "template": "classic"})
    assert res.status_code == 200
    assert res.mimetype == "application/pdf"
    assert res.data.startswith(b"%PDF")
    disposition = res.headers["Content-Disposition"]
    assert "attachment" in disposition
    assert "Nurul_Aisyah_Kamarudin_Resume.pdf" in disposition


@pytest.mark.parametrize("template", ["classic", "modern", "minimal"])
@pytest.mark.parametrize("size", ["A4", "Letter"])
# Dates "below" adds a line per entry, so it is left out: the full sample is meant to fill one page.
@pytest.mark.parametrize("style", [Style(), Style(accent="navy", body_align="justify", header_align="right", title_align="center")])
def test_sample_fits_one_page(template, size, style):
    html = render_html(parse_markdown(SAMPLE), template, size, style)
    doc = HTML(string=html, base_url=ROOT.as_uri() + "/", url_fetcher=LocalStaticFetcher()).render()
    assert len(doc.pages) == 1


def test_invalid_input_keeps_values(client):
    res = client.post("/pdf", data={"mode": "form", "headline": "Data Analyst", "phone": "+60 12-348 1927"})
    assert res.status_code == 400
    assert b"Please add your full name." in res.data
    assert b'value="Data Analyst"' in res.data
    assert b'value="+60 12-348 1927"' in res.data


def test_invalid_input_keeps_all_entries(client):
    data = MultiDict([("mode", "form"), ("experience_title", "Intern A"), ("experience_title", "Intern B")])
    res = client.post("/preview", data=data)
    assert res.status_code == 400
    assert b'value="Intern A"' in res.data and b'value="Intern B"' in res.data


def test_unknown_template_via_route(client):
    res = client.post("/pdf", data={"mode": "markdown", "markdown": SAMPLE, "template": "../app"})
    assert res.status_code == 400
    assert b"Please pick one of the templates." in res.data


def test_form_mode_pdf(client):
    res = client.post("/pdf", data={"mode": "form", "name": "Tan Wei Ming", "experience_title": "Intern", "experience_details": "Did things"})
    assert res.mimetype == "application/pdf"
    assert "Tan_Wei_Ming_Resume.pdf" in res.headers["Content-Disposition"]


def test_download_name():
    assert download_name("Nurul Aisyah binti Kamarudin", "pdf") == "Nurul_Aisyah_binti_Kamarudin_Resume.pdf"
    assert download_name("../../etc", "pdf") == "etc_Resume.pdf"
    assert download_name("", "docx") == "My_Resume.docx"
    assert download_name("陈伟明", "pdf") == "陈伟明_Resume.pdf"


def test_unicode_name_download_header(client):
    res = client.post("/pdf", data={"mode": "form", "name": "陈伟明", "skills_items": "Python"})
    assert "filename*=UTF-8''%E9%99%88%E4%BC%9F%E6%98%8E_Resume.pdf" in res.headers["Content-Disposition"]


def test_too_large(client):
    res = client.post("/pdf", data={"mode": "markdown", "markdown": "x" * (201 * 1024)})
    assert res.status_code == 413


# Preview pages and Word export

def test_page_images_match_pdf_pages():
    resume = parse_markdown(SAMPLE)
    one = render_page_images(resume, "classic")
    assert len(one) == 1 and one[0].data[:4] == b"RIFF"  # WebP
    assert one[0].height / one[0].width == pytest.approx(297 / 210, abs=0.002)
    letter = render_page_images(resume, "classic", "Letter")
    assert letter[0].height / letter[0].width == pytest.approx(11 / 8.5, abs=0.002)
    assert len(render_page_images(resume, "classic", style=Style(dates="below"))) == 2


def test_preview_warns_when_over_one_page(client):
    res = client.post("/preview", data={"mode": "markdown", "markdown": SAMPLE, "template": "classic", "dates": "below"})
    assert b"This runs to 2 pages" in res.data
    assert b"putting dates on the same line" in res.data
    assert b"Page 2 of 2" in res.data


def _docx_text(data: bytes):
    import io
    from docx import Document
    return Document(io.BytesIO(data))


def test_docx_download(client):
    res = client.post("/docx", data={"mode": "markdown", "markdown": SAMPLE, "template": "modern"})
    assert res.status_code == 200
    assert res.mimetype == DOCX_MIME
    assert "Nurul_Aisyah_Kamarudin_Resume.docx" in res.headers["Content-Disposition"]
    doc = _docx_text(res.data)
    text = [p.text for p in doc.paragraphs]
    assert text[0] == "Nurul Aisyah Kamarudin"
    assert "Experience" in text and "Certifications" in text
    # Title, a tab, then the date: right-aligned on the same line in Word.
    assert "Software Engineering Intern	Jun 2025 to Sep 2025" in text
    assert "Tanjung Data Sdn Bhd, Kuala Lumpur" in text
    bullets = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert any(p.text.startswith("Built a Flask service") for p in bullets)
    assert doc.core_properties.author == "Nurul Aisyah Kamarudin"


def test_docx_follows_style_and_page_size():
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches
    style = Style(accent="navy", header_align="right", body_align="justify", dates="below")
    doc = _docx_text(render_docx(parse_markdown(SAMPLE), "classic", "Letter", style))
    assert doc.sections[0].page_width == Inches(8.5)
    assert doc.paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert str(doc.paragraphs[0].runs[0].font.color.rgb) == "1F3A5F"
    head = next(p for p in doc.paragraphs if p.text.startswith("Software Engineering Intern"))
    assert "	" not in head.text and "Jun 2025" in head.text
    summary = next(p for p in doc.paragraphs if p.text.startswith("Computer Science graduate"))
    assert summary.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY


def test_docx_escapes_and_rejects_bad_template():
    r = from_form(MultiDict([("name", "A"), ("experience_title", "<script>x</script>"), ("skills_items", "**Python**")]))
    doc = _docx_text(render_docx(r, "minimal"))
    texts = [p.text for p in doc.paragraphs]
    assert "<script>x</script>" in texts  # shown as plain text, never markup
    skills = next(p for p in doc.paragraphs if p.text == "Python")
    assert skills.runs[0].font.bold
    with pytest.raises(ValueError):
        render_docx(r, "../app")
